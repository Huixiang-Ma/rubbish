"""知识摄取脚本：语料文件（CSV/HTML/TXT/MD，PDF/DOCX 可选）→ 切块 → 向量化 → document_chunks。

rag_lab 语料管道回填（tools/ingest_documents.py 的项目内版本）。
用法（backend 目录下）：
  python -m scripts.ingest_knowledge --dir ../data/ingest --tenant default
  python -m scripts.ingest_knowledge --file 景区政策.html --tenant t-001
"""
from __future__ import annotations

import argparse
import csv
import io
import re
import sys
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.services.chunker import chunk_text  # noqa: E402
from app.services.semantic import add_chunk  # noqa: E402

_BLOCK_TAGS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6",
               "section", "article", "table", "ul", "ol", "blockquote", "pre"}
_SKIP_TAGS = {"script", "style", "noscript", "template", "head", "svg"}


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip_depth = 0
        self._in_title = False
        self.title = ""
        self._parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: Any) -> None:
        if tag in _SKIP_TAGS:
            self._skip_depth += 1
        elif tag == "title":
            self._in_title = True
        elif tag in _BLOCK_TAGS:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIP_TAGS and self._skip_depth:
            self._skip_depth -= 1
        elif tag == "title":
            self._in_title = False
        elif tag in _BLOCK_TAGS:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._in_title:
            self.title += data.strip()
        if not self._skip_depth and data.strip():
            self._parts.append(data)


def parse_html(html: str) -> tuple[str, str]:
    parser = _TextExtractor()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        pass
    lines = [line.strip() for line in "".join(parser._parts).split("\n")]
    return parser.title.strip(), "\n".join(line for line in lines if line)


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def parse_csv_rows(path: Path) -> list[tuple[str, str]]:
    units: list[tuple[str, str]] = []
    for row in csv.DictReader(io.StringIO(read_text_file(path))):

        def get(*keys: str) -> str:
            return next((str(row[k]).strip() for k in keys if row.get(k) and str(row[k]).strip()), "")

        name = get("name", "名称", "景区名称", "景点名称")
        if not name:
            continue
        fragments = [name]
        city, address = get("city", "城市"), get("address", "地址")
        if city or address:
            fragments.append("位于" + " ".join(x for x in (city, address) if x))
        open_time = get("open_time", "开放时间", "营业时间")
        if open_time:
            fragments.append(f"开放时间：{open_time}")
        description = get("description", "简介")
        if description:
            fragments.append(description)
        text = "。".join(f.strip().rstrip("。") for f in fragments if f.strip()) + "。"
        text = "".join(text.split())[:400]
        if text:
            units.append((name, text))
    return units


def parse_document(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        return [parse_html(read_text_file(path))]
    if suffix == ".csv":
        return parse_csv_rows(path)
    if suffix in {".txt", ".md", ".markdown"}:
        return [(path.stem, read_text_file(path))]
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF 需要 pypdf：pip install pypdf") from exc
        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "").strip() for p in reader.pages]
        return [(path.stem, "\n".join(p for p in pages if p))]
    if suffix == ".docx":
        try:
            import docx
        except ImportError as exc:
            raise RuntimeError("DOCX 需要 python-docx：pip install python-docx") from exc
        document = docx.Document(str(path))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append("；".join(cells))
        return [(path.stem, "\n".join(parts))]
    raise ValueError(f"不支持的文档类型：{suffix}（支持 pdf/docx/html/htm/txt/md/csv）")


def main() -> None:
    parser = argparse.ArgumentParser(description="知识摄取：语料文件 → 切块 → 向量化 → document_chunks")
    parser.add_argument("--file", action="append", help="单个文件，可重复")
    parser.add_argument("--dir", action="append", help="目录批量（递归），可重复")
    parser.add_argument("--tenant", default="default", help="租户 ID（多租户隔离键）")
    parser.add_argument("--chunk-size", type=int, default=350)
    parser.add_argument("--chunk-overlap", type=int, default=70)
    args = parser.parse_args()

    targets: list[Path] = []
    for f in args.file or []:
        targets.append(Path(f))
    for d in args.dir or []:
        root = Path(d)
        allowed = {".pdf", ".docx", ".html", ".htm", ".txt", ".md", ".csv"}
        targets.extend(p for p in sorted(root.rglob("*")) if p.is_file() and p.suffix.lower() in allowed)
    if not targets:
        parser.error("至少提供 --file / --dir 之一")

    inserted = 0
    memory_mode = False
    for path in targets:
        try:
            units = parse_document(path)
        except Exception as exc:
            print(f"✗ {path.name}：{exc}")
            continue
        if not any(t.strip() for _, t in units):
            print(f"✗ {path.name}：无可提取文本（扫描件需 OCR），已跳过")
            continue
        file_chunks = 0
        for title, text in units:
            if not text.strip():
                continue
            chunks = chunk_text(text, size=args.chunk_size, overlap=args.chunk_overlap)
            chunks = [c for c in chunks if len(c) >= 30]
            chunks = [f"【{title}】{c}" for c in chunks]
            for chunk in chunks:
                out = add_chunk(f"{args.tenant}:{path.stem}", chunk, args.tenant)
                file_chunks += 1
                if out["mode"] == "memory":
                    memory_mode = True
        inserted += file_chunks
        print(f"✓ {path.name} → {file_chunks} 块（tenant={args.tenant}）")
    print(f"合计入库 {inserted} 块")
    if memory_mode:
        print("=" * 60)
        print("✗✗ 严重：部分块写入 MEMORY 模式（重启即失）！")
        print("  根因通常是 DATABASE_URL 在当前环境不可解析（宿主机跑脚本")
        print("  时 .env 里的 postgres:5432 是容器服务名）。")
        print("  修复：EMBEDDING/DATABASE 相关改用 localhost，或改走容器 API。")
        print("=" * 60)


if __name__ == "__main__":
    main()
